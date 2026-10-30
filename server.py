import io,json,random,secrets,time,threading,webbrowser,zipfile,os
from pathlib import Path
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from urllib.parse import urlparse
from email.parser import BytesParser
from email.policy import default
from docx import Document
from docx.shared import Cm
from docx.enum.section import WD_ORIENT
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
ROOT=Path(__file__).resolve().parent; PORT=int(os.environ.get('PORT','8773')); SESS={}; LOCK=threading.Lock()
SCENARIOS={
'word':[('Informe académico','Proyecto de investigación','Resultados preliminares','Universidad del Pacífico'),('Propuesta empresarial','Plan estratégico','Análisis financiero','Empresa Horizonte'),('Reporte ambiental','Desarrollo sostenible','Indicadores verdes','Instituto Ambiental'),('Manual de capacitación','Objetivos del curso','Plan de formación','Centro de Formación')],
'excel':[('Ventas regionales','Ventas'),('Inventario comercial','Inventario'),('Registro académico','Notas'),('Presupuesto institucional','Presupuesto')],
'ppt':[('Proyecto académico','Metodología'),('Informe de ventas','Resultados'),('Plan de formación','Actividades'),('Innovación digital','Objetivos')]}
BANK={
'word':[
('w_heading','Estilos','Aplica el estilo Título 1 al párrafo «{heading}».','Estilo Título 1 en el párrafo indicado.','Inicio > Estilos > Título 1.'),
('w_bold','Formato','Aplica negrita a la frase «{phrase}».','Frase completa en negrita.','Selecciona la frase y usa Ctrl+B.'),
('w_italic','Formato','Aplica cursiva a la frase «Texto de referencia».','Frase completa en cursiva.','Selecciona la frase y usa Ctrl+I.'),
('w_header','Encabezados','Agrega un encabezado con «{org}».','Texto indicado en el encabezado.','Insertar > Encabezado.'),
('w_footer','Pies de página','Agrega el texto «Documento de práctica» al pie de página.','Texto solicitado en el pie de página.','Insertar > Pie de página.'),
('w_table','Tablas','Inserta una tabla con 4 columnas y al menos 3 filas.','Tabla de cuatro columnas y tres filas como mínimo.','Insertar > Tabla > 4 × 3.'),
('w_landscape','Diseño','Cambia la orientación del documento a Horizontal.','Orientación horizontal.','Disposición > Orientación > Horizontal.'),
('w_margin','Diseño','Configura el margen superior en 2 cm.','Margen superior de 2 cm.','Disposición > Márgenes > Márgenes personalizados.'),
('w_paragraph','Edición','Inserta al final un párrafo con «Conclusión del proyecto».','Último párrafo con el texto exacto.','Ve al final y crea un nuevo párrafo.'),
('w_pagebreak','Saltos','Inserta un salto de página en el documento.','Al menos un salto de página explícito.','Insertar > Salto de página o Ctrl+Enter.'),
('w_highlight','Formato','Resalta en amarillo «Texto de referencia».','Resaltado amarillo de la frase.','Inicio > Color de resaltado del texto.'),
('w_alignment','Párrafos','Centra el párrafo «{heading}».','Alineación centrada en el párrafo.','Selecciona el párrafo y pulsa Ctrl+E.')],
'excel':[
('e_multiply','Fórmulas','En D2 escribe una fórmula que multiplique B2 por C2.','=B2*C2 en D2.','Selecciona D2 y escribe =B2*C2.'),
('e_fill','Referencias','Calcula D3:D7 multiplicando B por C en cada fila.','Fórmulas B3*C3 hasta B7*C7.','Copia la fórmula de D2 hacia abajo.'),
('e_sum','Funciones','En G2 calcula la suma de D2:D7.','=SUMA(D2:D7) o =SUM(D2:D7).','Usa la función SUMA.'),
('e_average','Funciones','En G3 calcula el promedio de C2:C7.','=PROMEDIO(C2:C7) o =AVERAGE(C2:C7).','Utiliza PROMEDIO.'),
('e_max','Funciones','En G4 calcula el máximo de C2:C7.','=MAX(C2:C7).','Usa MAX.'),
('e_min','Funciones','En G5 calcula el mínimo de C2:C7.','=MIN(C2:C7).','Usa MIN.'),
('e_if','Funciones','En E2 usa SI: si B2 es mayor que 4 devuelve «Alto»; si no, «Bajo».','=SI(B2>4,"Alto","Bajo").','Escribe =SI(B2>4;"Alto";"Bajo") según tu idioma de Office.'),
('e_currency','Formato','Aplica formato Moneda al intervalo C2:C7.','Formato de moneda en las seis celdas.','Selecciona C2:C7 > Formato de celdas > Moneda.'),
('e_freeze','Vista','Inmoviliza la fila superior.','Panel inmovilizado en A2.','Vista > Inmovilizar > Inmovilizar fila superior.'),
('e_chart','Gráficos','Inserta un gráfico de columnas de los datos del libro.','Un gráfico de columnas en la hoja.','Insertar > Gráfico de columnas.'),
('e_validation','Validación','Agrega validación de datos de número entero entre 1 y 100 en B2:B7.','Validación de entero entre 1 y 100.','Datos > Validación de datos.'),
('e_conditional','Formato condicional','Agrega una regla de formato condicional para valores mayores que 60 en C2:C7.','Regla condicional mayor que 60.','Inicio > Formato condicional > Reglas para resaltar celdas.')],
'ppt':[
('p_slide','Diapositivas','Agrega una cuarta diapositiva titulada «Próximos pasos».','Cuarta diapositiva titulada Próximos pasos.','Inicio > Nueva diapositiva.'),
('p_title','Edición','Cambia el título de la segunda diapositiva a «Indicadores clave».','Título exacto en la segunda diapositiva.','Edita el marcador de título.'),
('p_text','Cuadros de texto','Agrega «Evaluación del proyecto» en la tercera diapositiva.','Texto solicitado en la tercera diapositiva.','Insertar > Cuadro de texto.'),
('p_rect','Formas','Inserta un rectángulo en la segunda diapositiva.','Una autoforma rectangular en la diapositiva 2.','Insertar > Formas > Rectángulo.'),
('p_oval','Formas','Inserta un óvalo en la tercera diapositiva.','Una autoforma ovalada en la diapositiva 3.','Insertar > Formas > Óvalo.'),
('p_notes','Notas','Agrega «Revisar los resultados» a las notas del presentador.','Notas con el texto solicitado.','Abre el panel Notas.'),
('p_table','Tablas','Inserta una tabla de 3 filas y 3 columnas en la segunda diapositiva.','Tabla de al menos 3 × 3.','Insertar > Tabla > 3 × 3.'),
('p_chart','Gráficos','Inserta un gráfico en la tercera diapositiva.','Un objeto de gráfico en la diapositiva 3.','Insertar > Gráfico.'),
('p_hide','Presentación','Oculta la tercera diapositiva.','Tercera diapositiva marcada como oculta.','Clic derecho en diapositiva 3 > Ocultar diapositiva.'),
('p_fifth','Diapositivas','Agrega una quinta diapositiva titulada «Anexos».','Al menos cinco diapositivas y la quinta titulada Anexos.','Inicio > Nueva diapositiva.'),
('p_subtitle','Contenido','En la primera diapositiva agrega el subtítulo «Simulación MOS».','Texto en diapositiva 1.','Selecciona el marcador o agrega un cuadro de texto.'),
('p_speaker','Notas','En las notas escribe también «Presentación final».','Notas que incluyan Presentación final.','Escribe el texto en Notas.')]
}

def context(app):
 s=random.choice(SCENARIOS[app]);return ({'title':s[0],'heading':s[1],'phrase':s[2],'org':s[3]} if app=='word' else {'title':s[0],'sheet':s[1] if app=='excel' else s[1]})
def generate_doc(ctx,selected,solved=False):
 d=Document();d.add_heading(ctx['title'],0)
 p=d.add_paragraph(ctx['heading']);q=d.add_paragraph(ctx['phrase']);r=d.add_paragraph('Texto de referencia');d.add_paragraph('Los resultados serán utilizados para analizar el proyecto.')
 def has(k):return k in selected and solved
 if has('w_heading'):p.style='Heading 1'
 if has('w_alignment'):
  from docx.enum.text import WD_ALIGN_PARAGRAPH
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if has('w_bold'):q.runs[0].bold=True
 if has('w_italic'):r.runs[0].italic=True
 if has('w_highlight'):
  from docx.enum.text import WD_COLOR_INDEX
  r.runs[0].font.highlight_color=WD_COLOR_INDEX.YELLOW
 if has('w_header'):d.sections[0].header.paragraphs[0].text=ctx['org']
 if has('w_footer'):d.sections[0].footer.paragraphs[0].text='Documento de práctica'
 if has('w_table'):
  t=d.add_table(rows=3,cols=4)
  for i,row in enumerate(t.rows):
   for j,c in enumerate(row.cells):c.text=f'Dato {i+1}.{j+1}'
 if has('w_landscape'):
  sec=d.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
 if has('w_margin'):d.sections[0].top_margin=Cm(2)
 if has('w_pagebreak'):d.add_page_break()
 if has('w_paragraph'):d.add_paragraph('Conclusión del proyecto')
 f=io.BytesIO();d.save(f);return f.getvalue()
def generate_excel(ctx,selected,solved=False):
 w=Workbook();s=w.active;s.title=ctx['sheet'];s.append(['Concepto','Cantidad','Precio','Importe','Categoría',None,'Resumen'])
 for i in range(2,8):s.cell(i,1,f'Producto {i-1}');s.cell(i,2,i+1);s.cell(i,3,i*17)
 if solved:
  if 'e_multiply' in selected or 'e_fill' in selected:
   for i in range(2 if 'e_multiply' in selected else 3,8 if 'e_fill' in selected else 3):s.cell(i,4,f'=B{i}*C{i}')
  for k,addr,f in [('e_sum','G2','=SUM(D2:D7)'),('e_average','G3','=AVERAGE(C2:C7)'),('e_max','G4','=MAX(C2:C7)'),('e_min','G5','=MIN(C2:C7)'),('e_if','E2','=IF(B2>4,"Alto","Bajo")')]:
   if k in selected:s[addr]=f
  if 'e_currency' in selected:
   for i in range(2,8):s[f'C{i}'].number_format='"S/" #,##0.00'
  if 'e_freeze' in selected:s.freeze_panes='A2'
  if 'e_chart' in selected:
   c=BarChart();c.title='Comparación';c.add_data(Reference(s,min_col=3,min_row=1,max_row=7),titles_from_data=True);s.add_chart(c,'I3')
  if 'e_validation' in selected:
   v=DataValidation(type='whole',operator='between',formula1=1,formula2=100);s.add_data_validation(v);v.add('B2:B7')
  if 'e_conditional' in selected:s.conditional_formatting.add('C2:C7',CellIsRule(operator='greaterThan',formula=['60']))
 f=io.BytesIO();w.save(f);return f.getvalue()
def generate_ppt(ctx,selected,solved=False):
 p=Presentation()
 for name in [ctx['title'],ctx['sheet'],'Conclusiones']:
  slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text=name
 def has(k):return solved and k in selected
 if has('p_title'):p.slides[1].shapes.title.text='Indicadores clave'
 if has('p_subtitle'):p.slides[0].shapes.placeholders[1].text='Simulación MOS'
 if has('p_text'):
  sh=p.slides[2].shapes.add_textbox(Inches(1),Inches(2),Inches(5),Inches(1));sh.text='Evaluación del proyecto'
 if has('p_rect'):p.slides[1].shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(3),Inches(2),Inches(1))
 if has('p_oval'):p.slides[2].shapes.add_shape(MSO_SHAPE.OVAL,Inches(1),Inches(3),Inches(2),Inches(1))
 if has('p_notes'):p.slides[0].notes_slide.notes_text_frame.text='Revisar los resultados'
 if has('p_speaker'):p.slides[1].notes_slide.notes_text_frame.text='Presentación final'
 if has('p_table'):
  sh=p.slides[1].shapes.add_table(3,3,Inches(1),Inches(2),Inches(6),Inches(2))
  for row in sh.table.rows:
   for c in row.cells:c.text='Dato'
 if has('p_chart'):
  from pptx.chart.data import CategoryChartData
  from pptx.enum.chart import XL_CHART_TYPE
  data=CategoryChartData();data.categories=['A','B'];data.add_series('Serie 1',[4,8]);p.slides[2].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(2),Inches(5),Inches(3),data)
 if has('p_hide'):
  from pptx.oxml.xmlchemy import OxmlElement
  p.slides[2]._element.set('show','0')
 if has('p_slide'):
  s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Próximos pasos'
 if has('p_fifth'):
  while len(p.slides)<4:
   s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Complementos'
  s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Anexos'
 f=io.BytesIO();p.save(f);return f.getvalue()
BUILD={'word':generate_doc,'excel':generate_excel,'ppt':generate_ppt};EXT={'word':'docx','excel':'xlsx','ppt':'pptx'}
def grade(app,ctx,data):
 if app=='word':
  d=Document(io.BytesIO(data));ps=list(d.paragraphs)
  def line(t):return next((p for p in ps if t in p.text),None)
  p=line(ctx['heading']);q=line(ctx['phrase']);r=line('Texto de referencia')
  from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_COLOR_INDEX
  from docx.enum.section import WD_ORIENT
  check={'w_heading':bool(p and p.style.name.lower() in ('heading 1','título 1','titulo 1')),'w_bold':bool(q and any(ctx['phrase'] in run.text and run.bold for run in q.runs)),'w_italic':bool(r and any('Texto de referencia' in x.text and x.italic for x in r.runs)),'w_header':any(ctx['org'] in x.text for sec in d.sections for x in sec.header.paragraphs),'w_footer':any('Documento de práctica' in x.text for sec in d.sections for x in sec.footer.paragraphs),'w_table':any(len(t.rows)>=3 and len(t.columns)>=4 for t in d.tables),'w_landscape':d.sections[0].orientation==WD_ORIENT.LANDSCAPE,'w_margin':abs(d.sections[0].top_margin.cm-2)<0.12,'w_paragraph':bool(ps and ps[-1].text.strip()=='Conclusión del proyecto'),'w_pagebreak':any('w:type="page"' in run._r.xml or "w:type='page'" in run._r.xml for x in ps for run in x.runs),'w_highlight':bool(r and any('Texto de referencia' in x.text and x.font.highlight_color==WD_COLOR_INDEX.YELLOW for x in r.runs)),'w_alignment':bool(p and p.alignment==WD_ALIGN_PARAGRAPH.CENTER)}
 elif app=='excel':
  w=load_workbook(io.BytesIO(data));s=w[ctx['sheet']] if ctx['sheet'] in w else w.active
  f=lambda c:str(s[c].value or '').upper().replace(' ','').replace('$','')
  def mul(i):return f('D'+str(i)) in (f'=B{i}*C{i}',f'=C{i}*B{i}')
  def fun(c,*v):return f(c) in v
  check={'e_multiply':mul(2),'e_fill':all(mul(i) for i in range(3,8)),'e_sum':fun('G2','=SUM(D2:D7)','=SUMA(D2:D7)'),'e_average':fun('G3','=AVERAGE(C2:C7)','=PROMEDIO(C2:C7)'),'e_max':fun('G4','=MAX(C2:C7)'),'e_min':fun('G5','=MIN(C2:C7)'),'e_if':fun('E2','=IF(B2>4,"ALTO","BAJO")','=SI(B2>4,"ALTO","BAJO")'),'e_currency':all(any(t in s[f'C{i}'].number_format for t in ['S/','$','€','[$']) for i in range(2,8)),'e_freeze':s.freeze_panes=='A2','e_chart':any(c.__class__.__name__=='BarChart' for c in s._charts),'e_validation':any(v.type=='whole' and str(v.formula1)=='1' and str(v.formula2)=='100' and 'B2:B7' in str(v.sqref) for v in s.data_validations.dataValidation),'e_conditional':any('C2:C7' in str(k) for k in s.conditional_formatting)}
 else:
  p=Presentation(io.BytesIO(data));sl=list(p.slides);titles=[s.shapes.title.text.strip() if s.shapes.title else '' for s in sl]
  text=lambda i:' '.join(sh.text for sh in sl[i].shapes if sh.has_text_frame) if len(sl)>i else ''
  notes=' '.join(s.notes_slide.notes_text_frame.text for s in sl if s.has_notes_slide)
  check={'p_slide':len(sl)>=4 and 'Próximos pasos' in titles,'p_title':len(sl)>1 and titles[1]=='Indicadores clave','p_text':'Evaluación del proyecto' in text(2),'p_rect':len(sl)>1 and any(sh.is_placeholder is False and sh.shape_type==1 and sh.auto_shape_type==MSO_SHAPE.RECTANGLE for sh in sl[1].shapes),'p_oval':len(sl)>2 and any(sh.is_placeholder is False and sh.shape_type==1 and sh.auto_shape_type==MSO_SHAPE.OVAL for sh in sl[2].shapes),'p_notes':'Revisar los resultados' in notes,'p_table':len(sl)>1 and any(sh.has_table and len(sh.table.rows)>=3 and len(sh.table.columns)>=3 for sh in sl[1].shapes),'p_chart':len(sl)>2 and any(sh.has_chart for sh in sl[2].shapes),'p_hide':len(sl)>2 and sl[2]._element.get('show')=='0','p_fifth':len(sl)>=5 and titles[4]=='Anexos','p_subtitle':'Simulación MOS' in text(0),'p_speaker':'Presentación final' in notes}
 return check

from extra_tasks import EXTRA,apply as apply_extra,checks as extra_checks
from visual_tasks import NEW,apply as apply_visual,checks as visual_checks,asset_bytes,ASSET_NAME
for _app in BANK:
 BANK[_app].extend(EXTRA[_app]);BANK[_app].extend(NEW[_app])
_BASE_BUILD=BUILD.copy()
for _app in BUILD:
 def _make(ctx,selected,solved=False,app=_app):
  original=_BASE_BUILD[app](ctx,selected,solved)
  return apply_visual(app,apply_extra(app,original,selected,ctx),selected) if solved else original
 BUILD[_app]=_make
_base_grade=grade
def grade(app,ctx,data):
 results=_base_grade(app,ctx,data)
 results.update(extra_checks(app,data,ctx));results.update(visual_checks(app,data))
 return results

def new_exam(app):
 ctx=context(app)
 # Muestreo sin tareas contradictorias sobre el mismo elemento.
 groups=[{'w_justify','w_right'}, {'w_bullets','w_numbering'}, {'w_landscape','w_page_size'}, {'p_sixth','p_seventh'}, {'w_spacing','w_line_15','w_line_exact'}, {'w_image_insert','w_image_width'}, {'e_image_insert','e_image_width'}, {'p_image_insert','p_image_width'}]
 candidates=BANK[app][:];random.shuffle(candidates);sel=[];areas=set()
 def fits(item):
  key=item[0]
  return all(not (key in group and any(chosen in group for chosen in sel)) for group in groups)
 for item in candidates:
  if len(sel)>=5:break
  if item[1] not in areas and fits(item):sel.append(item[0]);areas.add(item[1])
 for item in candidates:
  if len(sel)>=5:break
  if item[0] not in sel and fits(item):sel.append(item[0])
 token=secrets.token_urlsafe(12)
 with LOCK:SESS[token]={'app':app,'ctx':ctx,'selected':sel,'created':time.time(),'graded':False}
 definitions={x[0]:x for x in BANK[app]};tasks=[{'id':id,'area':definitions[id][1],'text':definitions[id][2].format(**ctx),'points':200} for id in sel]
 return {'id':token,'app':app,'title':ctx['title'],'tasks':tasks,'file':f'MOS_{app}_{token[:6]}.{EXT[app]}','created':SESS[token]['created']}
class H(BaseHTTPRequestHandler):
 def send(self,code,obj,typ='application/json',filename=None):
  b=obj if isinstance(obj,bytes) else json.dumps(obj,ensure_ascii=False).encode();self.send_response(code);self.send_header('Content-Type',typ);self.send_header('Content-Length',str(len(b)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff')
  self.send_header('X-Frame-Options','DENY')
  self.send_header('Referrer-Policy','no-referrer')
  if filename:self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
  self.end_headers();self.wfile.write(b)
 def do_GET(self):
  path=urlparse(self.path).path
  if path=='/health':return self.send(200,{'status':'ok','version':'3.3'})
  if path=='/api/resource/image':return self.send(200,asset_bytes(),'image/png',ASSET_NAME)
  if path=='/':return self.send(200,(ROOT/'index.html').read_bytes(),'text/html; charset=utf-8')
  if path.startswith('/api/file/'):
   _,api,file_segment,kind,token=path.split('/')
   info=SESS.get(token)
   if not info or kind not in ('base','solution'):return self.send(404,{'error':'Examen no encontrado'})
   if kind=='solution' and not info.get('graded'):return self.send(403,{'error':'La solución se habilita después de entregar la evaluación'})
   app=info['app'];data=BUILD[app](info['ctx'],info['selected'],kind=='solution');return self.send(200,data,'application/octet-stream',f'{kind}_{app}_{token[:6]}.{EXT[app]}')
  return self.send(404,{'error':'No encontrado'})
 def do_POST(self):
  path=urlparse(self.path).path;n=int(self.headers.get('Content-Length','0'))
  if n>17000000:return self.send(413,{'error':'Máximo 15 MB'})
  if n>17000000:return self.send(413,{'error':'Máximo 15 MB'})
  raw=self.rfile.read(n)
  try:
   if path=='/api/new':
    app=json.loads(raw or b'{}')['app'];assert app in BANK
    with LOCK:
     for key,exam in list(SESS.items()):
      if time.time()-exam['created']>10800:SESS.pop(key,None)
    return self.send(200,new_exam(app))
   if path=='/api/grade':
    typ=self.headers.get('Content-Type','');msg=BytesParser(policy=default).parsebytes(('Content-Type: '+typ+'\r\nMIME-Version: 1.0\r\n\r\n').encode()+raw);fields={x.get_param('name',header='content-disposition'):x for x in msg.iter_parts()};token=fields['id'].get_content().strip();info=SESS.get(token)
    if not info:raise ValueError('Examen inexistente. Genera uno nuevo.')
    uploaded=fields['file'];name=uploaded.get_filename() or '';app=info['app']
    if not name.lower().endswith('.'+EXT[app]):raise ValueError('Archivo de Office de tipo incorrecto.')
    data=uploaded.get_payload(decode=True)
    if len(data)>15000000:raise ValueError('Archivo demasiado grande.')
    checks=grade(app,info['ctx'],data);defs={x[0]:x for x in BANK[app]};res=[]
    for id in info['selected']:
     d=defs[id];res.append({'title':d[2].format(**info['ctx']),'area':d[1],'success':bool(checks[id]),'points':200 if checks[id] else 0,'expected':d[3],'correction':d[4]})
    with LOCK:info['graded']=True
    score=sum(x['points'] for x in res)
    return self.send(200,{'score':score,'max':1000,'items':res,'solution':'/api/file/solution/'+token})
   raise ValueError('Ruta desconocida')
  except Exception as e:self.send(400,{'error':str(e)})
 def log_message(self,*args):pass
if __name__=='__main__':
 host='0.0.0.0' if os.environ.get('PORT') else '127.0.0.1'
 s=ThreadingHTTPServer((host,PORT),H);print('MOS Office Simulator 3.3 en %s:%s'%(host,PORT),flush=True)
 if host=='127.0.0.1' and '--no-browser' not in __import__('sys').argv:threading.Timer(1,lambda:webbrowser.open('http://127.0.0.1:%s'%PORT)).start()
 try:s.serve_forever()
 except KeyboardInterrupt:pass
