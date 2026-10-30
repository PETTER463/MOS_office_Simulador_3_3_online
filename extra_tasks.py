"""20 tareas adicionales para cada aplicación; generador y evaluador de propiedades OOXML."""
import io,re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.chart import LineChart,PieChart,Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches,Pt as PPTPt
from pptx.dml.color import RGBColor as PPTRGB
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData

EXTRA={
'word':[
('w_font16','Tipografía','Establece en 16 puntos el tamaño de letra del párrafo «{heading}».','Tamaño de 16 puntos aplicado al texto del párrafo.','Selecciona el párrafo y elige 16 en Inicio > Fuente.'),
('w_fontcolor','Color de fuente','Cambia a rojo el color de la frase «{phrase}».','Texto de la frase en rojo RGB FF0000.','Selecciona la frase y aplica color rojo en Inicio > Color de fuente.'),
('w_underline','Subrayado','Subraya la frase «{phrase}».','La frase completa aparece subrayada.','Selecciona la frase y pulsa Ctrl+S.'),
('w_justify','Alineación','Justifica el párrafo «Texto de referencia».','Párrafo justificado.','Inicio > Párrafo > Justificar (Ctrl+J).'),
('w_right','Alineación','Alinea a la derecha el párrafo «Texto de referencia».','Párrafo alineado a la derecha.','Inicio > Párrafo > Alinear a la derecha.'),
('w_spacing','Párrafos','Establece interlineado doble en «Texto de referencia».','Interlineado de valor 2.0.','Inicio > Párrafo > Opciones de interlineado > Doble.'),
('w_before','Espaciado','Configura 12 puntos de espacio antes de «Texto de referencia».','Espacio anterior de 12 puntos.','Párrafo > Espaciado > Anterior 12 pt.'),
('w_after','Espaciado','Configura 18 puntos de espacio después de «Texto de referencia».','Espacio posterior de 18 puntos.','Párrafo > Espaciado > Posterior 18 pt.'),
('w_indent','Sangrías','Aplica sangría izquierda de 1 cm al párrafo «Texto de referencia».','Sangría izquierda de 1 cm.','Párrafo > Sangría > Izquierda 1 cm.'),
('w_firstline','Sangrías','Aplica sangría de primera línea de 1 cm al párrafo «Texto de referencia».','Sangría de primera línea de 1 cm.','Párrafo > Sangría especial > Primera línea.'),
('w_bullets','Listas','Convierte la frase «{phrase}» en un elemento de lista con viñetas.','Párrafo con estilo Lista con viñetas o List Bullet.','Selecciona el párrafo y utiliza Inicio > Viñetas.'),
('w_numbering','Listas','Convierte la frase «{phrase}» en un elemento de lista numerada.','Párrafo con estilo Lista numerada o List Number.','Selecciona el párrafo y elige Numeración.'),
('w_title2','Estilos','Aplica el estilo Título 2 al párrafo «Texto de referencia».','Estilo Título 2.','Inicio > Estilos > Título 2.'),
('w_table_style','Tablas','Inserta una tabla 2 × 2 y aplícale el estilo Table Grid (Tabla con cuadrícula).','Tabla 2 × 2 con el estilo Table Grid.','Insertar > Tabla y aplicar estilo de cuadrícula.'),
('w_table_center','Tablas','Inserta una tabla 2 × 3 y céntrala horizontalmente.','Tabla de 2 filas y 3 columnas centrada.','Insertar > Tabla y Propiedades de tabla > Alineación: centro.'),
('w_page_size','Diseño','Cambia el tamaño de página a A4.','Papel A4: 21 × 29,7 cm.','Disposición > Tamaño > A4.'),
('w_bottom_margin','Diseño','Configura el margen inferior en 1,5 cm.','Margen inferior 1,5 cm.','Disposición > Márgenes > Márgenes personalizados.'),
('w_right_margin','Diseño','Configura el margen derecho en 3 cm.','Margen derecho de 3 cm.','Disposición > Márgenes > Márgenes personalizados.'),
('w_core_title','Propiedades','Establece el título del documento en «{title} - Edición final».','Propiedad Título actualizada.','Archivo > Información > Propiedades > Título.'),
('w_core_subject','Propiedades','Establece el asunto del documento en «Certificación MOS».','Propiedad Asunto actualizada.','Archivo > Información > Propiedades > Asunto.')],
'excel':[
('e_count','Funciones','En G6 calcula cuántas celdas numéricas hay en B2:B7 con CONTAR.','Fórmula =COUNT(B2:B7).','Escribe =CONTAR(B2:B7) o =COUNT(B2:B7) en G6.'),
('e_counta','Funciones','En G7 calcula las celdas no vacías de A2:A7 con CONTARA.','Fórmula =COUNTA(A2:A7).','Escribe =CONTARA(A2:A7) o =COUNTA(A2:A7).'),
('e_countif','Funciones','En G8 cuenta los precios mayores que 60 de C2:C7.','Fórmula =COUNTIF(C2:C7,">60").','Utiliza CONTAR.SI con el criterio >60.'),
('e_sumif','Funciones','En G9 suma los precios de C2:C7 solo cuando B2:B7 sea mayor que 4.','Fórmula =SUMIF(B2:B7,">4",C2:C7).','Utiliza SUMAR.SI con rango B2:B7 y suma C2:C7.'),
('e_median','Estadística','En G10 calcula la mediana de C2:C7.','Fórmula =MEDIAN(C2:C7).','Usa MEDIANA.'),
('e_round','Redondeo','En G11 redondea el valor de C2 a cero decimales.','Fórmula =ROUND(C2,0).','Utiliza REDONDEAR(C2;0).'),
('e_abs','Funciones','En G12 calcula el valor absoluto de B2 menos 10.','Fórmula =ABS(B2-10).','Usa ABS(B2-10).'),
('e_concat','Texto','En G13 concatena A2, un espacio y A3.','Fórmula que concatene A2, un espacio y A3.','Utiliza =A2&" "&A3.'),
('e_upper','Texto','En G14 convierte a mayúsculas el contenido de A2.','Fórmula =UPPER(A2).','Utiliza MAYUSC(A2).'),
('e_len','Texto','En G15 calcula el número de caracteres de A2.','Fórmula =LEN(A2).','Utiliza LARGO(A2).'),
('e_percent','Formato','Aplica formato Porcentaje con dos decimales a F2:F7.','Las seis celdas tienen formato 0.00%.','Inicio > Número > Porcentaje; dos decimales.'),
('e_boldhead','Formato','Aplica negrita a los encabezados A1:G1.','Encabezados en negrita.','Selecciona A1:G1 y pulsa Ctrl+N.'),
('e_bluehead','Formato','Aplica fondo azul a las celdas A1:G1.','Relleno azul RGB 4472C4 en los encabezados.','Inicio > Color de relleno > Más colores.'),
('e_centerhead','Alineación','Centra horizontalmente las celdas A1:G1.','Encabezados centrados.','Inicio > Alineación > Centrar.'),
('e_colwidth','Columnas','Establece el ancho de la columna A en 24.','Ancho de columna A = 24.','Inicio > Formato > Ancho de columna.'),
('e_rowheight','Filas','Establece la altura de la fila 1 en 30 puntos.','Altura de fila 1 = 30.','Inicio > Formato > Alto de fila.'),
('e_tabcolor','Hojas','Aplica color rojo a la pestaña de la hoja.','Color rojo RGB FF0000 en la pestaña.','Clic derecho en pestaña > Color de etiqueta.'),
('e_printarea','Impresión','Define A1:G7 como área de impresión.','Área de impresión $A$1:$G$7.','Diseño de página > Área de impresión > Establecer.'),
('e_linechart','Gráficos','Inserta un gráfico de líneas en la hoja.','Existe un gráfico de líneas.','Insertar > Gráfico de líneas.'),
('e_piechart','Gráficos','Inserta un gráfico circular en la hoja.','Existe un gráfico circular.','Insertar > Gráfico circular.')],
'ppt':[
('p_sixth','Diapositivas','Agrega una sexta diapositiva titulada «Resumen ejecutivo».','Sexta diapositiva titulada Resumen ejecutivo.','Inicio > Nueva diapositiva; escribe el título.'),
('p_slide3title','Edición','Cambia el título de la tercera diapositiva a «Análisis de resultados».','Título exacto de la diapositiva 3.','Selecciona la tercera diapositiva y modifica el título.'),
('p_roundrect','Formas','Inserta un rectángulo redondeado en la primera diapositiva.','Autoforma de rectángulo redondeado en diapositiva 1.','Insertar > Formas > Rectángulo redondeado.'),
('p_triangle','Formas','Inserta un triángulo isósceles en la segunda diapositiva.','Triángulo isósceles en diapositiva 2.','Insertar > Formas > Triángulo.'),
('p_chevron','Formas','Inserta un chevrón en la tercera diapositiva.','Forma de chevrón en diapositiva 3.','Insertar > Formas > Chevrón.'),
('p_arrow','Formas','Inserta una flecha hacia la derecha en la segunda diapositiva.','Flecha derecha en diapositiva 2.','Insertar > Formas > Flecha derecha.'),
('p_diamond','Formas','Inserta un rombo en la primera diapositiva.','Rombo en diapositiva 1.','Insertar > Formas > Rombo.'),
('p_star','Formas','Inserta una estrella de cinco puntas en la tercera diapositiva.','Estrella de cinco puntas en diapositiva 3.','Insertar > Formas > Estrella de cinco puntas.'),
('p_line','Formas','Inserta una línea recta en la segunda diapositiva.','Línea recta en diapositiva 2.','Insertar > Formas > Línea.'),
('p_piechart','Gráficos','Inserta un gráfico circular en la primera diapositiva.','Gráfico circular en diapositiva 1.','Insertar > Gráfico > Circular.'),
('p_linechart','Gráficos','Inserta un gráfico de líneas en la segunda diapositiva.','Gráfico de líneas en diapositiva 2.','Insertar > Gráfico > Línea.'),
('p_table4','Tablas','Inserta una tabla de 4 filas × 2 columnas en la tercera diapositiva.','Tabla de 4 filas y 2 columnas.','Insertar > Tabla > 4 filas × 2 columnas.'),
('p_textred','Color de texto','Agrega el texto «Prioridad alta» en rojo en la primera diapositiva.','Texto rojo Prioridad alta en diapositiva 1.','Insertar > Cuadro de texto; aplica color rojo.'),
('p_textbold','Tipografía','Agrega el texto «OBJETIVOS» en negrita en la segunda diapositiva.','Texto OBJETIVOS en negrita.','Insertar > Texto; aplica negrita.'),
('p_textcenter','Alineación','Agrega el texto «PLAN DE ACCIÓN» centrado en la tercera diapositiva.','Texto centrado en un cuadro de texto.','Selecciona el texto y céntralo.'),
('p_core_title','Propiedades','Establece el título de la presentación en «{title} - Final».','Propiedad Título actualizada.','Archivo > Información > Propiedades > Título.'),
('p_core_subject','Propiedades','Establece el asunto de la presentación en «Capacitación MOS».','Propiedad Asunto actualizada.','Archivo > Información > Propiedades > Asunto.'),
('p_slidewidth','Diseño','Configura el ancho de diapositiva en 13,333 pulgadas (pantalla ancha).','Ancho de diapositiva aproximadamente 13,333 pulgadas.','Diseño > Tamaño de diapositiva > Pantalla panorámica.'),
('p_notes3','Notas','Escribe «Conclusión principal» en las notas de la tercera diapositiva.','Nota solicitada en diapositiva 3.','Abre el panel Notas de la tercera diapositiva.'),
('p_seventh','Diapositivas','Agrega una séptima diapositiva titulada «Bibliografía».','Séptima diapositiva titulada Bibliografía.','Inicio > Nueva diapositiva; agrega el título.')]
}

def apply(app,data,selected,ctx):
 ids=set(selected)
 if not any(i in ids for i,*_ in EXTRA[app]):return data
 b=io.BytesIO()
 if app=='word':
  d=Document(io.BytesIO(data));ps=d.paragraphs;p=next(x for x in ps if ctx['heading'] in x.text);q=next(x for x in ps if ctx['phrase'] in x.text);r=next(x for x in ps if 'Texto de referencia' in x.text)
  if 'w_font16' in ids:p.runs[0].font.size=Pt(16)
  if 'w_fontcolor' in ids:q.runs[0].font.color.rgb=RGBColor(255,0,0)
  if 'w_underline' in ids:q.runs[0].underline=True
  if 'w_justify' in ids:r.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if 'w_right' in ids:r.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  if 'w_spacing' in ids:r.paragraph_format.line_spacing=2.0
  if 'w_before' in ids:r.paragraph_format.space_before=Pt(12)
  if 'w_after' in ids:r.paragraph_format.space_after=Pt(18)
  if 'w_indent' in ids:r.paragraph_format.left_indent=Pt(28.3465)
  if 'w_firstline' in ids:r.paragraph_format.first_line_indent=Pt(28.3465)
  if 'w_bullets' in ids:q.style='List Bullet'
  if 'w_numbering' in ids:q.style='List Number'
  if 'w_title2' in ids:r.style='Heading 2'
  if 'w_table_style' in ids:
   t=d.add_table(rows=2,cols=2);t.style='Table Grid'
  if 'w_table_center' in ids:
   t=d.add_table(rows=2,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER
  if 'w_page_size' in ids:
   from docx.shared import Cm
   d.sections[0].page_width=Cm(21);d.sections[0].page_height=Cm(29.7)
  if 'w_bottom_margin' in ids:
   from docx.shared import Cm
   d.sections[0].bottom_margin=Cm(1.5)
  if 'w_right_margin' in ids:
   from docx.shared import Cm
   d.sections[0].right_margin=Cm(3)
  if 'w_core_title' in ids:d.core_properties.title=ctx['title']+' - Edición final'
  if 'w_core_subject' in ids:d.core_properties.subject='Certificación MOS'
  d.save(b)
 elif app=='excel':
  w=load_workbook(io.BytesIO(data));s=w.active
  formulas={'e_count':('G6','=COUNT(B2:B7)'),'e_counta':('G7','=COUNTA(A2:A7)'),'e_countif':('G8','=COUNTIF(C2:C7,">60")'),'e_sumif':('G9','=SUMIF(B2:B7,">4",C2:C7)'),'e_median':('G10','=MEDIAN(C2:C7)'),'e_round':('G11','=ROUND(C2,0)'),'e_abs':('G12','=ABS(B2-10)'),'e_concat':('G13','=A2&" "&A3'),'e_upper':('G14','=UPPER(A2)'),'e_len':('G15','=LEN(A2)')}
  for k,(cell,v) in formulas.items():
   if k in ids:s[cell]=v
  if 'e_percent' in ids:
   for row in s['F2:F7']:
    for c in row:c.number_format='0.00%'
  if 'e_boldhead' in ids:
   for c in s[1][:7]:c.font=Font(bold=True)
  if 'e_bluehead' in ids:
   for c in s[1][:7]:c.fill=PatternFill('solid',fgColor='4472C4')
  if 'e_centerhead' in ids:
   for c in s[1][:7]:c.alignment=Alignment(horizontal='center')
  if 'e_colwidth' in ids:s.column_dimensions['A'].width=24
  if 'e_rowheight' in ids:s.row_dimensions[1].height=30
  if 'e_tabcolor' in ids:s.sheet_properties.tabColor='FF0000'
  if 'e_printarea' in ids:s.print_area='A1:G7'
  if 'e_linechart' in ids:
   chart=LineChart();chart.add_data(Reference(s,min_col=3,min_row=1,max_row=7),titles_from_data=True);s.add_chart(chart,'I10')
  if 'e_piechart' in ids:
   chart=PieChart();chart.add_data(Reference(s,min_col=3,min_row=1,max_row=7),titles_from_data=True);chart.set_categories(Reference(s,min_col=1,min_row=2,max_row=7));s.add_chart(chart,'I20')
  w.save(b)
 else:
  p=Presentation(io.BytesIO(data))
  def ensure(n):
   while len(p.slides)<n:p.slides.add_slide(p.slide_layouts[1])
  if 'p_sixth' in ids:ensure(6);p.slides[5].shapes.title.text='Resumen ejecutivo'
  if 'p_seventh' in ids:ensure(7);p.slides[6].shapes.title.text='Bibliografía'
  if 'p_slide3title' in ids:p.slides[2].shapes.title.text='Análisis de resultados'
  shapes={'p_roundrect':(0,MSO_SHAPE.ROUNDED_RECTANGLE),'p_triangle':(1,MSO_SHAPE.ISOSCELES_TRIANGLE),'p_chevron':(2,MSO_SHAPE.CHEVRON),'p_arrow':(1,MSO_SHAPE.RIGHT_ARROW),'p_diamond':(0,MSO_SHAPE.DIAMOND),'p_star':(2,MSO_SHAPE.STAR_5_POINT)}
  for k,(idx,typ) in shapes.items():
   if k in ids:p.slides[idx].shapes.add_shape(typ,Inches(1),Inches(2),Inches(2),Inches(1))
  if 'p_line' in ids:
   from pptx.enum.shapes import MSO_CONNECTOR
   p.slides[1].shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(1),Inches(1),Inches(3),Inches(2))
  for key,idx,typ in [('p_piechart',0,XL_CHART_TYPE.PIE),('p_linechart',1,XL_CHART_TYPE.LINE)]:
   if key in ids:
    chart=CategoryChartData();chart.categories=['A','B'];chart.add_series('Serie',[2,6]);p.slides[idx].shapes.add_chart(typ,Inches(1),Inches(2),Inches(5),Inches(3),chart)
  if 'p_table4' in ids:p.slides[2].shapes.add_table(4,2,Inches(1),Inches(2),Inches(5),Inches(2))
  for key,idx,txt in [('p_textred',0,'Prioridad alta'),('p_textbold',1,'OBJETIVOS'),('p_textcenter',2,'PLAN DE ACCIÓN')]:
   if key in ids:
    shape=p.slides[idx].shapes.add_textbox(Inches(1),Inches(4),Inches(5),Inches(1));shape.text=txt;par=shape.text_frame.paragraphs[0]
    if key=='p_textred':par.runs[0].font.color.rgb=PPTRGB(255,0,0)
    if key=='p_textbold':par.runs[0].font.bold=True
    if key=='p_textcenter':par.alignment=PP_ALIGN.CENTER
  if 'p_core_title' in ids:p.core_properties.title=ctx['title']+' - Final'
  if 'p_core_subject' in ids:p.core_properties.subject='Capacitación MOS'
  if 'p_slidewidth' in ids:p.slide_width=Inches(13.333)
  if 'p_notes3' in ids:p.slides[2].notes_slide.notes_text_frame.text='Conclusión principal'
  p.save(b)
 return b.getvalue()

def checks(app,data,ctx):
 res={}
 if app=='word':
  d=Document(io.BytesIO(data));p=next((p for p in d.paragraphs if ctx['heading'] in p.text),None);q=next((p for p in d.paragraphs if ctx['phrase'] in p.text),None);r=next((p for p in d.paragraphs if 'Texto de referencia' in p.text),None)
  def font(par,attribute):return bool(par and any(getattr(run.font,attribute,None) for run in par.runs))
  def near(v,x,tol=.2):return bool(v is not None and abs(v-x)<tol)
  res.update(w_font16=bool(p and any(run.font.size and near(run.font.size.pt,16) for run in p.runs)),w_fontcolor=bool(q and any(run.font.color.rgb==RGBColor(255,0,0) for run in q.runs)),w_underline=font(q,'underline'),w_justify=bool(r and r.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY),w_right=bool(r and r.alignment==WD_ALIGN_PARAGRAPH.RIGHT),w_spacing=bool(r and near(r.paragraph_format.line_spacing,2)),w_before=bool(r and r.paragraph_format.space_before and near(r.paragraph_format.space_before.pt,12)),w_after=bool(r and r.paragraph_format.space_after and near(r.paragraph_format.space_after.pt,18)),w_indent=bool(r and r.paragraph_format.left_indent and near(r.paragraph_format.left_indent.cm,1)),w_firstline=bool(r and r.paragraph_format.first_line_indent and near(r.paragraph_format.first_line_indent.cm,1)),w_bullets=bool(q and 'bullet' in q.style.name.lower()),w_numbering=bool(q and 'number' in q.style.name.lower()),w_title2=bool(r and r.style.name.lower() in ('heading 2','título 2','titulo 2')),w_table_style=any(len(t.rows)==2 and len(t.columns)==2 and t.style and t.style.name=='Table Grid' for t in d.tables),w_table_center=any(len(t.rows)==2 and len(t.columns)==3 and t.alignment==WD_TABLE_ALIGNMENT.CENTER for t in d.tables),w_page_size=near(d.sections[0].page_width.cm,21) and near(d.sections[0].page_height.cm,29.7),w_bottom_margin=near(d.sections[0].bottom_margin.cm,1.5),w_right_margin=near(d.sections[0].right_margin.cm,3),w_core_title=d.core_properties.title==ctx['title']+' - Edición final',w_core_subject=d.core_properties.subject=='Certificación MOS')
 elif app=='excel':
  w=load_workbook(io.BytesIO(data));s=w.active
  formulas={'e_count':('G6',['=COUNT(B2:B7)','=CONTAR(B2:B7)']),'e_counta':('G7',['=COUNTA(A2:A7)','=CONTARA(A2:A7)']),'e_countif':('G8',['=COUNTIF(C2:C7,">60")','=CONTAR.SI(C2:C7,">60")']),'e_sumif':('G9',['=SUMIF(B2:B7,">4",C2:C7)','=SUMAR.SI(B2:B7,">4",C2:C7)']),'e_median':('G10',['=MEDIAN(C2:C7)','=MEDIANA(C2:C7)']),'e_round':('G11',['=ROUND(C2,0)','=REDONDEAR(C2,0)']),'e_abs':('G12',['=ABS(B2-10)']),'e_concat':('G13',['=A2&" "&A3']),'e_upper':('G14',['=UPPER(A2)','=MAYUSC(A2)']),'e_len':('G15',['=LEN(A2)','=LARGO(A2)'])}
  for k,(c,expected) in formulas.items():res[k]=str(s[c].value or '').upper().replace(' ','').replace(';',',') in [v.upper().replace(' ','') for v in expected]
  res.update(e_percent=all('%' in s[f'F{i}'].number_format and '0.00' in s[f'F{i}'].number_format for i in range(2,8)),e_boldhead=all(c.font.bold for c in s[1][:7]),e_bluehead=all(c.fill.fgColor.rgb and str(c.fill.fgColor.rgb).upper().endswith('4472C4') for c in s[1][:7]),e_centerhead=all(c.alignment.horizontal=='center' for c in s[1][:7]),e_colwidth=abs((s.column_dimensions['A'].width or 0)-24)<.2,e_rowheight=abs((s.row_dimensions[1].height or 0)-30)<.2,e_tabcolor=bool(s.sheet_properties.tabColor and str(s.sheet_properties.tabColor.rgb).upper().endswith('FF0000')),e_printarea='A1:G7' in str(s.print_area).replace('$',''),e_linechart=any(isinstance(c,LineChart) for c in s._charts),e_piechart=any(isinstance(c,PieChart) for c in s._charts))
 else:
  p=Presentation(io.BytesIO(data));sl=list(p.slides)
  def shapes(idx,typ):return any(sh.shape_type==1 and sh.auto_shape_type==typ for sh in sl[idx].shapes)
  types={'p_roundrect':(0,MSO_SHAPE.ROUNDED_RECTANGLE),'p_triangle':(1,MSO_SHAPE.ISOSCELES_TRIANGLE),'p_chevron':(2,MSO_SHAPE.CHEVRON),'p_arrow':(1,MSO_SHAPE.RIGHT_ARROW),'p_diamond':(0,MSO_SHAPE.DIAMOND),'p_star':(2,MSO_SHAPE.STAR_5_POINT)}
  for k,(idx,typ) in types.items():res[k]=shapes(idx,typ)
  from pptx.enum.shapes import MSO_SHAPE_TYPE
  res['p_line']=any(sh.shape_type==MSO_SHAPE_TYPE.LINE for sh in sl[1].shapes)
  for key,idx,typename in [('p_piechart',0,'PieChart'),('p_linechart',1,'LineChart')]:res[key]=any(sh.has_chart and sh.chart.chart_type== (XL_CHART_TYPE.PIE if key=='p_piechart' else XL_CHART_TYPE.LINE) for sh in sl[idx].shapes)
  res['p_table4']=any(sh.has_table and len(sh.table.rows)==4 and len(sh.table.columns)==2 for sh in sl[2].shapes)
  def runs(idx,txt):return [(sh,pa,ru) for sh in sl[idx].shapes if sh.has_text_frame for pa in sh.text_frame.paragraphs for ru in pa.runs if txt in ru.text]
  res.update(p_textred=any(ru.font.color and ru.font.color.type and ru.font.color.rgb==PPTRGB(255,0,0) for sh,pa,ru in runs(0,'Prioridad alta')),p_textbold=any(ru.font.bold for sh,pa,ru in runs(1,'OBJETIVOS')),p_textcenter=any(pa.alignment==PP_ALIGN.CENTER for sh,pa,ru in runs(2,'PLAN DE ACCIÓN')),p_core_title=p.core_properties.title==ctx['title']+' - Final',p_core_subject=p.core_properties.subject=='Capacitación MOS',p_slidewidth=abs(p.slide_width/Inches(1)-13.333)<.02,p_notes3='Conclusión principal' in (sl[2].notes_slide.notes_text_frame.text if sl[2].has_notes_slide else ''),p_slide3title=sl[2].shapes.title.text=='Análisis de resultados',p_sixth=len(sl)>=6 and sl[5].shapes.title and sl[5].shapes.title.text=='Resumen ejecutivo',p_seventh=len(sl)>=7 and sl[6].shapes.title and sl[6].shapes.title.text=='Bibliografía')
 return res
