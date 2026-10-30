"""20 original practical tasks. Images are provided in an independent PNG asset."""
import io
from docx import Document
from docx.shared import Inches,Cm,Pt
from docx.enum.text import WD_LINE_SPACING,WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from pptx import Presentation
from pptx.util import Inches as PPTInches
from pptx.dml.color import RGBColor
from PIL import Image,ImageDraw
ASSET_NAME='MOS_Imagen_Practica.png'

def asset_bytes():
 im=Image.new('RGB',(640,360),'#ecf4fb');draw=ImageDraw.Draw(im)
 draw.rectangle((50,55,590,300),fill='#0e6397');draw.ellipse((100,95,240,235),fill='#f5c857');draw.rectangle((300,120,530,230),fill='#f7fafc')
 draw.text((305,145),'MOS 2026',fill='#173a58')
 out=io.BytesIO();im.save(out,format='PNG');return out.getvalue()

NEW={
'word':[
('w_image_insert','Imágenes','Inserta la imagen de práctica en el documento.','Documento con una imagen incrustada.','Descarga Imagen de práctica e insértala desde Insertar > Imágenes.'),
('w_image_width','Tamaño de imagen','Inserta la imagen de práctica y ajusta su ancho a 7 cm.','Una imagen con ancho de 7 cm (tolerancia 0,2 cm).','Insertar > Imágenes > Este dispositivo; Formato de imagen > Ancho 7 cm.'),
('w_page_number','Numeración de páginas','Inserta un campo automático de número de página en el pie de página.','Pie de página con campo PAGE, no un número escrito manualmente.','Insertar > Número de página > Final de página.'),
('w_line_15','Interlineado','Aplica interlineado de 1,5 líneas al párrafo «Texto de referencia».','Interlineado de 1,5 en el párrafo indicado.','Inicio > Párrafo > Interlineado > 1,5.'),
('w_line_exact','Interlineado','Establece interlineado Exacto de 18 puntos en «Texto de referencia».','Interlineado exacto de 18 puntos.','Párrafo > Interlineado: Exacto; En: 18 pt.'),
('w_header_first','Encabezados','Activa «Primera página diferente» para el encabezado.','Propiedad Primera página diferente activada.','Doble clic en encabezado > Primera página diferente.'),
('w_columns','Columnas','Organiza el documento en dos columnas.','Configuración de dos columnas para la sección.','Disposición > Columnas > Dos.'),
('w_hyperlink','Hipervínculos','Inserta un hipervínculo a https://www.microsoft.com dentro del documento.','Enlace funcional a www.microsoft.com.','Insertar > Vínculo; introduce https://www.microsoft.com.'),
('w_keep_next','Paginación de párrafos','Activa «Conservar con el siguiente» en «{heading}».','Propiedad mantener con el siguiente activada.','Párrafo > Líneas y saltos de página > Conservar con el siguiente.'),
('w_keep_together','Paginación de párrafos','Activa «Conservar líneas juntas» en «Texto de referencia».','Propiedad mantener líneas juntas activada.','Párrafo > Líneas y saltos de página > Conservar líneas juntas.'),
('w_table_shading','Formato de tablas','Inserta una tabla de 2 filas × 2 columnas y aplica sombreado amarillo a la primera celda.','Tabla 2 × 2 con celda superior izquierda amarilla.','Insertar > Tabla; selecciona primera celda > Diseño de tabla > Sombreado amarillo.'),
('w_repeat_header','Tablas','Inserta una tabla de 3 filas × 2 columnas y marca la primera fila como fila de encabezado repetida.','La primera fila tiene la propiedad repetir encabezado.','Propiedades de tabla > Fila > Repetir como fila de encabezado.')],
'excel':[
('e_image_insert','Imágenes','Inserta la imagen de práctica en la hoja.','Hoja con al menos una imagen incrustada.','Insertar > Ilustraciones > Imágenes.'),
('e_image_width','Tamaño de imagen','Inserta la imagen de práctica y cambia su ancho a 300 píxeles aproximadamente.','Imagen incrustada con ancho de aproximadamente 300 px.','Selecciona la imagen y ajusta su tamaño desde Formato de imagen.'),
('e_page_number','Configuración de página','Inserta el campo automático número de página en el pie de página central.','Pie de página central con el código &P.','Diseño de página > Configurar página > Encabezado y pie de página.'),
('e_landscape','Configuración de página','Cambia la orientación de la hoja a horizontal.','Orientación horizontal en la hoja.','Diseño de página > Orientación > Horizontal.')],
'ppt':[
('p_image_insert','Imágenes','Inserta la imagen de práctica en la segunda diapositiva.','Al menos una imagen incrustada en la diapositiva 2.','Insertar > Imágenes > Este dispositivo.'),
('p_image_width','Tamaño de imagen','Inserta la imagen de práctica en la primera diapositiva y dale 3 pulgadas de ancho (7,62 cm).','Imagen de 7,62 cm de ancho, con tolerancia.','Formato de imagen > Tamaño > Ancho: 7,62 cm.'),
('p_slide_number','Numeración','Inserta un número de diapositiva visible en la tercera diapositiva.','Campo de número de diapositiva en la diapositiva 3.','Insertar > Encabezado y pie de página > Número de diapositiva > Aplicar.'),
('p_slide_background','Fondo','Cambia el fondo de la primera diapositiva a azul RGB 4472C4.','Color de fondo azul #4472C4 en la diapositiva 1.','Diseño > Formato del fondo > Relleno sólido > RGB 68,114,196.')]
}

def element(tag):return OxmlElement(tag)
def fld(run,name):
 b=element('w:fldChar');b.set(qn('w:fldCharType'),'begin');run._r.append(b)
 t=element('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+name+' ';run._r.append(t)
 e=element('w:fldChar');e.set(qn('w:fldCharType'),'end');run._r.append(e)

def apply(app,data,selected):
 chosen=set(selected);f=io.BytesIO()
 if app=='word':
  d=Document(io.BytesIO(data));ps=d.paragraphs;heading=ps[1];reference=next(p for p in ps if 'Texto de referencia' in p.text)
  if 'w_image_insert' in chosen:d.add_picture(io.BytesIO(asset_bytes()),width=Cm(5))
  if 'w_image_width' in chosen:d.add_picture(io.BytesIO(asset_bytes()),width=Cm(7))
  if 'w_page_number' in chosen:fld(d.sections[0].footer.paragraphs[0].add_run(),'PAGE')
  if 'w_line_15' in chosen:reference.paragraph_format.line_spacing=1.5
  if 'w_line_exact' in chosen:reference.paragraph_format.line_spacing=Pt(18)
  if 'w_header_first' in chosen:d.sections[0].different_first_page_header_footer=True
  if 'w_columns' in chosen:
   sect=d.sections[0]._sectPr;cols=sect.xpath('./w:cols')
   if cols:cols[0].set(qn('w:num'),'2')
   else:
    c=element('w:cols');c.set(qn('w:num'),'2');sect.append(c)
  if 'w_hyperlink' in chosen:
   from docx.opc.constants import RELATIONSHIP_TYPE as RT
   p=d.add_paragraph('Sitio Microsoft: ');rid=p.part.relate_to('https://www.microsoft.com',RT.HYPERLINK,is_external=True);h=element('w:hyperlink');h.set(qn('r:id'),rid);r=element('w:r');t=element('w:t');t.text='Microsoft';r.append(t);h.append(r);p._p.append(h)
  if 'w_keep_next' in chosen:heading.paragraph_format.keep_with_next=True
  if 'w_keep_together' in chosen:reference.paragraph_format.keep_together=True
  if 'w_table_shading' in chosen:
   t=d.add_table(rows=2,cols=2);tcPr=t.cell(0,0)._tc.get_or_add_tcPr();sh=element('w:shd');sh.set(qn('w:fill'),'FFFF00');tcPr.append(sh)
  if 'w_repeat_header' in chosen:
   t=d.add_table(rows=3,cols=2);trPr=t.rows[0]._tr.get_or_add_trPr();trPr.append(element('w:tblHeader'))
  if 'w_paragraph' in chosen and d.paragraphs[-1].text.strip()!='Conclusión del proyecto':d.add_paragraph('Conclusión del proyecto')
  d.save(f)
 elif app=='excel':
  w=load_workbook(io.BytesIO(data));s=w.active
  if 'e_image_insert' in chosen or 'e_image_width' in chosen:
   image=XLImage(io.BytesIO(asset_bytes()));image.width=300 if 'e_image_width' in chosen else 250;image.height=169 if 'e_image_width' in chosen else 141;s.add_image(image,'I12')
  if 'e_page_number' in chosen:s.oddFooter.center.text='&P'
  if 'e_landscape' in chosen:s.page_setup.orientation='landscape'
  w.save(f)
 else:
  p=Presentation(io.BytesIO(data));slides=p.slides
  if 'p_image_insert' in chosen:slides[1].shapes.add_picture(io.BytesIO(asset_bytes()),PPTInches(1),PPTInches(2),width=PPTInches(2.5))
  if 'p_image_width' in chosen:slides[0].shapes.add_picture(io.BytesIO(asset_bytes()),PPTInches(1),PPTInches(2),width=PPTInches(3))
  if 'p_slide_number' in chosen:
   # Office slide number placeholder text field
   from pptx.oxml.xmlchemy import OxmlElement as PE
   s=slides[2];sh=s.shapes.add_textbox(PPTInches(8),PPTInches(6.5),PPTInches(.6),PPTInches(.3));tf=sh.text_frame;tf.clear();para=tf.paragraphs[0];fldel=PE('a:fld');fldel.set('id','{00000000-0000-0000-0000-000000000001}');fldel.set('type','slidenum');rPr=PE('a:rPr');fldel.append(rPr);tx=PE('a:t');tx.text='3';fldel.append(tx);para._p.append(fldel)
  if 'p_slide_background' in chosen:slides[0].background.fill.solid();slides[0].background.fill.fore_color.rgb=RGBColor(68,114,196)
  p.save(f)
 return f.getvalue()

def checks(app,data):
 out={}
 if app=='word':
  d=Document(io.BytesIO(data));r=next((p for p in d.paragraphs if 'Texto de referencia' in p.text),None);heading=d.paragraphs[1] if len(d.paragraphs)>1 else None
  inline=list(d.inline_shapes)
  out['w_image_insert']=len(inline)>=1
  out['w_image_width']=any(abs(s.width.cm-7)<.2 for s in inline)
  out['w_page_number']=any('PAGE' in sec.footer._element.xml for sec in d.sections)
  out['w_line_15']=bool(r and r.paragraph_format.line_spacing is not None and abs(float(r.paragraph_format.line_spacing)-1.5)<.03)
  out['w_line_exact']=bool(r and r.paragraph_format.line_spacing is not None and hasattr(r.paragraph_format.line_spacing,'pt') and abs(r.paragraph_format.line_spacing.pt-18)<.4)
  out['w_header_first']=any(sec.different_first_page_header_footer for sec in d.sections)
  out['w_columns']=any('w:num="2"' in sec._sectPr.xml and '<w:cols' in sec._sectPr.xml for sec in d.sections)
  out['w_hyperlink']=any('https://www.microsoft.com' in rel.target_ref for rel in d.part.rels.values() if rel.reltype.endswith('/hyperlink'))
  out['w_keep_next']=bool(heading and heading.paragraph_format.keep_with_next)
  out['w_keep_together']=bool(r and r.paragraph_format.keep_together)
  out['w_table_shading']=any(len(t.rows)>=2 and len(t.columns)>=2 and 'FFFF00' in t.cell(0,0)._tc.xml for t in d.tables)
  out['w_repeat_header']=any(len(t.rows)>=3 and len(t.columns)>=2 and 'tblHeader' in t.rows[0]._tr.xml for t in d.tables)
 elif app=='excel':
  w=load_workbook(io.BytesIO(data));s=w.active;imgs=s._images
  out['e_image_insert']=len(imgs)>=1
  out['e_image_width']=any(abs(getattr(getattr(i.anchor,'ext',None),'cx',0)/9525-300)<8 for i in imgs)
  out['e_page_number']='&P' in (s.oddFooter.center.text or '')
  out['e_landscape']=s.page_setup.orientation=='landscape'
 else:
  p=Presentation(io.BytesIO(data));sl=p.slides
  out['p_image_insert']=any(sh.shape_type==13 for sh in sl[1].shapes)
  out['p_image_width']=any(sh.shape_type==13 and abs(sh.width.inches-3)<.1 for sh in sl[0].shapes)
  out['p_slide_number']='type="slidenum"' in sl[2]._element.xml or 'type="sldNum"' in sl[2]._element.xml
  fill=sl[0].background.fill
  out['p_slide_background']=False
  if fill.type is not None:
   try:out['p_slide_background']=str(fill.fore_color.rgb).upper()=='4472C4'
   except (TypeError,AttributeError):pass
 return out
